import os
import shutil
import uuid

# =============================================================================
# LAZY IMPORTS FRAMEWORK
# We import these inside functions to ensure the server starts even if libs are missing.
# =============================================================================

def convert_pdf_to_docx(input_path: str, output_path: str):
    """
    Converts PDF to Word (DOCX) using pdf2docx.
    """
    try:
        from pdf2docx import Converter
    except ImportError:
        raise Exception("pdf2docx library not installed.")
        
    try:
        cv = Converter(input_path)
        cv.convert(output_path, start=0, end=None)
        cv.close()
    except Exception as e:
        raise Exception(f"PDF to DOCX conversion failed: {str(e)}")

def convert_pdf_to_pptx(input_path: str, output_path: str):
    """
    Converts PDF to PowerPoint (PPTX).
    Strategy: Convert pages to images, then place on slides to preserve diagrams.
    """
    try:
        import fitz  # PyMuPDF
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        raise Exception("PyMuPDF or python-pptx library not installed.")

    try:
        doc = fitz.open(input_path)
        prs = Presentation()
        
        # Standard Slide Layout (Blank)
        BLANK_SLIDE_LAYOUT_INDEX = 6 
        
        for i, page in enumerate(doc):
            # 1. Render page to high-res image
            zoom = 2  # 2x zoom = ~144dpi, good balance
            mat = fitz.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=mat)
            
            # Save temp image
            img_path = f"{input_path}_temp_page_{i}.png"
            pix.save(img_path)
            
            # 2. Create Slide
            slide = prs.slides.add_slide(prs.slide_layouts[BLANK_SLIDE_LAYOUT_INDEX])
            
            # 3. Add image to slide (Fit to slide)
            # Default PPT size is 10x7.5 inches. 
            left = top = Inches(0)
            height = Inches(7.5) # Standard height
            
            slide.shapes.add_picture(img_path, left, top, height=height)
            
            # Cleanup temp image
            if os.path.exists(img_path):
                os.remove(img_path)
        
        prs.save(output_path)
        
    except Exception as e:
        raise Exception(f"PDF to PPTX conversion failed: {str(e)}")

